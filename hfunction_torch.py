"""CPU/PyTorch implementations of the h-function interval refinement.

The project workbooks store interval endpoints on the Poisson mean scale
theta=lambda+b. The module keeps endpoints on that scale to preserve the
workbook's numerical convention. Search candidates are signal rates lambda.
A 1e-12 tolerance in the test-statistic comparison keeps mathematical ties
stable under floating-point arithmetic.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import numpy as np
import torch
from openpyxl import load_workbook
from scipy.stats import poisson


Backend = Literal["cpu", "torch"]
EXACT_INITIAL_METHODS = {"FC", "EB"}
# The h-function uses <=, so absorb only double-precision roundoff at ties.
STATISTIC_TIE_TOL = 1e-12


@dataclass(frozen=True)
class IntervalTable:
    """Dense interval table indexed by observed count y, on the mean scale theta."""

    lower: np.ndarray
    upper: np.ndarray

    def __post_init__(self) -> None:
        lower = np.asarray(self.lower, dtype=np.float64)
        upper = np.asarray(self.upper, dtype=np.float64)
        if lower.ndim != 1 or upper.ndim != 1 or len(lower) != len(upper):
            raise ValueError("lower and upper must be same-length one-dimensional arrays")
        if len(lower) == 0:
            raise ValueError("an interval table cannot be empty")
        if not np.all(np.isfinite(lower)) or not np.all(np.isfinite(upper)):
            raise ValueError("interval endpoints must be finite")
        if np.any(lower < -1e-10) or np.any(upper < lower - 1e-10):
            raise ValueError("intervals must satisfy 0 <= lower <= upper")
        object.__setattr__(self, "lower", np.maximum(lower, 0.0))
        object.__setattr__(self, "upper", np.maximum(upper, 0.0))

    @property
    def max_y(self) -> int:
        return len(self.lower) - 1

    def require_y(self, y: np.ndarray | list[int]) -> np.ndarray:
        y = np.asarray(y, dtype=np.int64)
        if y.ndim != 1 or np.any(y < 0) or np.any(y > self.max_y):
            raise ValueError(f"interval table covers y=0..{self.max_y}; requested values are outside it")
        return y


def load_interval_sheet(
    path: str | Path,
    sheet_name: str = "Initial Interval",
    background: float = 3.0,
    source_scale: Literal["mean", "signal"] = "mean",
    max_y: int | None = None,
) -> IntervalTable:
    """Load columns (y, lower, upper) from a project workbook.

    The supplied project workbooks store the first three columns on the mean
    scale. Rows must cover consecutive integer counts starting at zero.
    """

    path = Path(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise KeyError(f"sheet {sheet_name!r} not found in {path.name}")
        rows: dict[int, tuple[float, float]] = {}
        for row in workbook[sheet_name].iter_rows(values_only=True):
            if len(row) < 3 or any(value is None for value in row[:3]):
                continue
            try:
                y = int(row[0])
                lower = float(row[1])
                upper = float(row[2])
            except (TypeError, ValueError):
                continue
            if float(row[0]) != y:
                raise ValueError(f"noninteger y value {row[0]!r} in {path.name}/{sheet_name}")
            if max_y is not None and y > max_y:
                continue
            if y in rows:
                raise ValueError(f"duplicate y={y} in {path.name}/{sheet_name}")
            rows[y] = (lower, upper)
    finally:
        workbook.close()

    if not rows:
        raise ValueError(f"no interval rows found in {path.name}/{sheet_name}")
    target_max_y = max(rows) if max_y is None else max_y
    expected = set(range(target_max_y + 1))
    if set(rows) != expected:
        missing = sorted(expected - set(rows))[:10]
        raise ValueError(f"nonconsecutive y values in {path.name}/{sheet_name}; missing {missing}")

    lower = np.array([rows[y][0] for y in range(target_max_y + 1)], dtype=np.float64)
    upper = np.array([rows[y][1] for y in range(target_max_y + 1)], dtype=np.float64)
    if source_scale == "signal":
        lower = lower + background
        upper = upper + background
    elif source_scale != "mean":
        raise ValueError("source_scale must be 'mean' or 'signal'")
    return IntervalTable(lower=lower, upper=upper)


def _validate_inputs(
    observed_y: np.ndarray | list[int],
    signal_grid: np.ndarray | list[float],
    table: IntervalTable,
    background: float,
    quantile: float,
) -> tuple[np.ndarray, np.ndarray]:
    observed_y = table.require_y(observed_y)
    signal_grid = np.asarray(signal_grid, dtype=np.float64)
    if signal_grid.ndim != 1 or len(signal_grid) == 0:
        raise ValueError("signal_grid must be a nonempty one-dimensional array")
    if np.any(signal_grid < 0) or np.any(np.diff(signal_grid) < 0):
        raise ValueError("signal_grid must be nonnegative and increasing")
    if background <= 0:
        raise ValueError("background must be positive")
    if not 0.0 < quantile < 1.0:
        raise ValueError("quantile must lie strictly between zero and one")
    return observed_y, signal_grid


def _poisson_cutoffs(signal_grid: np.ndarray, background: float, quantile: float) -> np.ndarray:
    """Use SciPy for inverse CDFs so CPU and GPU use identical integer supports."""

    cutoffs = np.asarray(poisson.ppf(quantile, signal_grid + background), dtype=np.float64)
    if not np.all(np.isfinite(cutoffs)):
        raise ValueError("could not calculate finite Poisson quantiles")
    return cutoffs.astype(np.int64)


_POISSON_CUTOFF_LUTS: dict[tuple[float, float, float], np.ndarray] = {}


def _grid_indices(signal_values: np.ndarray, grid_step: float) -> np.ndarray | None:
    """Return exact grid indices when every signal value lies on the search grid."""

    if grid_step <= 0:
        return None
    indices = np.rint(signal_values / grid_step).astype(np.int64)
    if np.any(indices < 0):
        return None
    expected = np.round(indices * grid_step, 10)
    if not np.array_equal(signal_values, expected):
        return None
    return indices


def _poisson_cutoff_lut(
    max_grid_index: int,
    background: float,
    quantile: float,
    grid_step: float,
) -> np.ndarray:
    """Cache SciPy cutoffs once per q and grid point, preserving its exact definition."""

    key = (float(background), float(quantile), float(grid_step))
    cached = _POISSON_CUTOFF_LUTS.get(key)
    start = 0 if cached is None else len(cached)
    if start <= max_grid_index:
        indices = np.arange(start, max_grid_index + 1, dtype=np.float64)
        signal_values = np.round(indices * grid_step, 10)
        new_cutoffs = _poisson_cutoffs(signal_values, background, quantile)
        cached = new_cutoffs if cached is None else np.concatenate((cached, new_cutoffs))
        _POISSON_CUTOFF_LUTS[key] = cached
    return cached


def h_pairs_cpu(
    observed_y: np.ndarray | list[int],
    signal_values: np.ndarray | list[float],
    table: IntervalTable,
    background: float = 3.0,
    quantile: float = 0.999999,
    pair_chunk: int = 4096,
    grid_step: float | None = None,
) -> np.ndarray:
    """Evaluate h for paired (y, lambda) candidates without a Cartesian grid."""

    observed_y = table.require_y(observed_y)
    signal_values = np.asarray(signal_values, dtype=np.float64)
    if signal_values.ndim != 1 or signal_values.shape != observed_y.shape:
        raise ValueError("signal_values and observed_y must be same-length vectors")
    output = np.empty(len(signal_values), dtype=np.float64)
    if len(signal_values) == 0:
        return output
    indices = _grid_indices(signal_values, grid_step) if grid_step is not None else None
    if indices is not None:
        cutoff_lut = _poisson_cutoff_lut(
            int(indices.max()), background, quantile, grid_step
        )
    else:
        cutoff_lut = None
        cutoffs_all = _poisson_cutoffs(signal_values, background, quantile)
    for start in range(0, len(signal_values), pair_chunk):
        stop = min(start + pair_chunk, len(signal_values))
        lambdas = signal_values[start:stop]
        ys = observed_y[start:stop]
        cutoffs = (
            cutoff_lut[indices[start:stop]]
            if indices is not None
            else cutoffs_all[start:stop]
        )
        max_k = int(cutoffs.max())
        if max_k > table.max_y:
            raise ValueError(f"h-function needs y={max_k}, but interval table ends at y={table.max_y}")
        k = np.arange(max_k + 1, dtype=np.float64)
        theta = lambdas + background
        pmf = poisson.pmf(k[:, None], theta[None, :])
        sample_t = np.minimum(
            theta[None, :] - table.lower[: max_k + 1, None],
            table.upper[: max_k + 1, None] - theta[None, :],
        )
        observed_t = np.minimum(
            theta - table.lower[ys],
            table.upper[ys] - theta,
        )
        extreme = (sample_t <= observed_t[None, :] + STATISTIC_TIE_TOL) & (
            k[:, None] <= cutoffs[None, :]
        )
        output[start:stop] = np.sum(pmf * extreme, axis=0, dtype=np.float64)
    return output


def h_pairs_torch(
    observed_y: np.ndarray | list[int],
    signal_values: np.ndarray | list[float],
    table: IntervalTable,
    background: float = 3.0,
    quantile: float = 0.999999,
    device: str | torch.device | None = None,
    pair_chunk: int = 16384,
    grid_step: float | None = None,
) -> np.ndarray:
    """Evaluate paired h values with PyTorch float64 on CPU or CUDA."""

    observed_y = table.require_y(observed_y)
    signal_values = np.asarray(signal_values, dtype=np.float64)
    if signal_values.ndim != 1 or signal_values.shape != observed_y.shape:
        raise ValueError("signal_values and observed_y must be same-length vectors")
    if pair_chunk < 1:
        raise ValueError("pair_chunk must be positive")
    if device is None:
        device = "cuda" if torch.cuda.is_available() else "cpu"
    device = torch.device(device)
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA was requested but PyTorch reports no available CUDA device")

    output = np.empty(len(signal_values), dtype=np.float64)
    if len(signal_values) == 0:
        return output
    indices = _grid_indices(signal_values, grid_step) if grid_step is not None else None
    if indices is not None:
        cutoff_lut = _poisson_cutoff_lut(
            int(indices.max()), background, quantile, grid_step
        )
        max_cutoff = int(cutoff_lut[int(indices.max())])
        cutoffs_all = None
    else:
        cutoff_lut = None
        cutoffs_all = _poisson_cutoffs(signal_values, background, quantile)
        max_cutoff = int(cutoffs_all.max())
    if max_cutoff > table.max_y:
        raise ValueError(f"h-function needs y={max_cutoff}, but interval table ends at y={table.max_y}")

    # These values are shared across chunks. Keeping them on the device avoids
    # repeated host-to-device copies and recomputing the factorial term.
    k_int_all = torch.arange(max_cutoff + 1, dtype=torch.int64, device=device)
    k_float_all = k_int_all.to(torch.float64)
    log_factorial_all = torch.lgamma(k_float_all + 1.0)
    # The sample-space endpoints are needed only through max_cutoff, but the
    # observed statistic also indexes the interval table at observed_y, which
    # can exceed that cutoff on a full (y, lambda) grid.
    table_max_y = max(max_cutoff, int(observed_y.max()))
    lower_table = torch.as_tensor(table.lower[: table_max_y + 1], dtype=torch.float64, device=device)
    upper_table = torch.as_tensor(table.upper[: table_max_y + 1], dtype=torch.float64, device=device)

    with torch.inference_mode():
        for start in range(0, len(signal_values), pair_chunk):
            stop = min(start + pair_chunk, len(signal_values))
            lambdas_np = signal_values[start:stop]
            ys_np = observed_y[start:stop]
            cutoffs_np = (
                cutoff_lut[indices[start:stop]]
                if indices is not None
                else cutoffs_all[start:stop]
            )
            max_k = int(cutoffs_np.max())

            lambdas = torch.as_tensor(lambdas_np, dtype=torch.float64, device=device)
            theta = lambdas + background
            k = k_float_all[: max_k + 1]
            k_int = k_int_all[: max_k + 1]
            log_factorial = log_factorial_all[: max_k + 1]
            cutoffs = torch.as_tensor(cutoffs_np, dtype=torch.int64, device=device)
            lower_k = lower_table[: max_k + 1]
            upper_k = upper_table[: max_k + 1]
            ys = torch.as_tensor(ys_np, dtype=torch.int64, device=device)
            lower_y = lower_table[ys]
            upper_y = upper_table[ys]
            log_pmf = (
                k[:, None] * torch.log(theta[None, :])
                - theta[None, :]
                - log_factorial[:, None]
            )
            pmf = torch.exp(log_pmf)
            sample_t = torch.minimum(
                theta[None, :] - lower_k[:, None],
                upper_k[:, None] - theta[None, :],
            )
            observed_t = torch.minimum(theta - lower_y, upper_y - theta)
            extreme = (sample_t <= observed_t[None, :] + STATISTIC_TIE_TOL) & (
                k_int[:, None] <= cutoffs[None, :]
            )
            values = torch.sum(pmf * extreme, dim=0, dtype=torch.float64)
            output[start:stop] = values.cpu().numpy()
    return output


def h_values_cpu(
    observed_y: np.ndarray | list[int],
    signal_grid: np.ndarray | list[float],
    table: IntervalTable,
    background: float = 3.0,
    quantile: float = 0.999999,
    lambda_chunk: int = 128,
    y_chunk: int = 32,
) -> np.ndarray:
    """Return h(y; lambda) with shape (len(signal_grid), len(observed_y))."""

    observed_y, signal_grid = _validate_inputs(observed_y, signal_grid, table, background, quantile)
    output = np.empty((len(signal_grid), len(observed_y)), dtype=np.float64)

    for start in range(0, len(signal_grid), lambda_chunk):
        stop = min(start + lambda_chunk, len(signal_grid))
        lambdas = signal_grid[start:stop]
        cutoffs = _poisson_cutoffs(lambdas, background, quantile)
        max_k = int(cutoffs.max())
        if max_k > table.max_y:
            raise ValueError(f"h-function needs y={max_k}, but interval table ends at y={table.max_y}")
        k = np.arange(max_k + 1, dtype=np.float64)
        theta = lambdas + background
        pmf = poisson.pmf(k[:, None], theta[None, :])
        sample_t = np.minimum(
            theta[None, :] - table.lower[: max_k + 1, None],
            table.upper[: max_k + 1, None] - theta[None, :],
        )
        retained = k[:, None] <= cutoffs[None, :]

        for y_start in range(0, len(observed_y), y_chunk):
            y_stop = min(y_start + y_chunk, len(observed_y))
            ys = observed_y[y_start:y_stop]
            observed_t = np.minimum(
                theta[None, :] - table.lower[ys, None],
                table.upper[ys, None] - theta[None, :],
            ).T
            extreme = (
                sample_t[:, :, None] <= observed_t[None, :, :] + STATISTIC_TIE_TOL
            ) & retained[:, :, None]
            output[start:stop, y_start:y_stop] = np.sum(
                pmf[:, :, None] * extreme,
                axis=0,
                dtype=np.float64,
            )
    return output


def h_values_torch(
    observed_y: np.ndarray | list[int],
    signal_grid: np.ndarray | list[float],
    table: IntervalTable,
    background: float = 3.0,
    quantile: float = 0.999999,
    device: str | torch.device | None = None,
    lambda_chunk: int = 128,
    y_chunk: int = 32,
) -> np.ndarray:
    """Return h(y; lambda) using PyTorch float64, optionally on a CUDA device."""

    observed_y, signal_grid = _validate_inputs(observed_y, signal_grid, table, background, quantile)
    if device is None:
        device = "cuda" if torch.cuda.is_available() else "cpu"
    device = torch.device(device)
    if device.type == "cuda" and not torch.cuda.is_available():
        raise RuntimeError("CUDA was requested but PyTorch reports no available CUDA device")

    output = np.empty((len(signal_grid), len(observed_y)), dtype=np.float64)
    with torch.inference_mode():
        for start in range(0, len(signal_grid), lambda_chunk):
            stop = min(start + lambda_chunk, len(signal_grid))
            lambdas_np = signal_grid[start:stop]
            cutoffs_np = _poisson_cutoffs(lambdas_np, background, quantile)
            max_k = int(cutoffs_np.max())
            if max_k > table.max_y:
                raise ValueError(f"h-function needs y={max_k}, but interval table ends at y={table.max_y}")

            lambdas = torch.as_tensor(lambdas_np, dtype=torch.float64, device=device)
            theta = lambdas + background
            k = torch.arange(max_k + 1, dtype=torch.float64, device=device)
            cutoffs = torch.as_tensor(cutoffs_np, dtype=torch.int64, device=device)
            lower_k = torch.as_tensor(table.lower[: max_k + 1], dtype=torch.float64, device=device)
            upper_k = torch.as_tensor(table.upper[: max_k + 1], dtype=torch.float64, device=device)
            log_pmf = (
                k[:, None] * torch.log(theta[None, :])
                - theta[None, :]
                - torch.lgamma(k[:, None] + 1.0)
            )
            pmf = torch.exp(log_pmf)
            sample_t = torch.minimum(
                theta[None, :] - lower_k[:, None],
                upper_k[:, None] - theta[None, :],
            )
            retained = k.to(torch.int64)[:, None] <= cutoffs[None, :]

            for y_start in range(0, len(observed_y), y_chunk):
                y_stop = min(y_start + y_chunk, len(observed_y))
                ys_np = observed_y[y_start:y_stop]
                lower_y = torch.as_tensor(table.lower[ys_np], dtype=torch.float64, device=device)
                upper_y = torch.as_tensor(table.upper[ys_np], dtype=torch.float64, device=device)
                observed_t = torch.minimum(
                    theta[None, :] - lower_y[:, None],
                    upper_y[:, None] - theta[None, :],
                ).T
                extreme = (
                    sample_t[:, :, None] <= observed_t[None, :, :] + STATISTIC_TIE_TOL
                ) & retained[:, :, None]
                values = torch.sum(pmf[:, :, None] * extreme, dim=0, dtype=torch.float64)
                output[start:stop, y_start:y_stop] = values.cpu().numpy()
    return output


def _grid_from_upper(upper: np.ndarray, step: float) -> np.ndarray:
    if step <= 0:
        raise ValueError("grid step must be positive")
    max_upper = float(np.max(upper))
    last_index = int(np.floor(max_upper / step + 1e-10))
    return np.round(np.arange(last_index + 1, dtype=np.float64) * step, 10)


def refine_intervals(
    table: IntervalTable,
    observed_y: np.ndarray | list[int],
    search_lower: np.ndarray,
    search_upper: np.ndarray,
    alpha: float = 0.1,
    quantile: float = 0.999999,
    background: float = 3.0,
    step: float = 0.01,
    backend: Backend = "torch",
    device: str | torch.device | None = None,
    pair_chunk: int = 16384,
) -> IntervalTable:
    """Refine selected observations by taking the hull of grid points with h>alpha."""

    observed_y = table.require_y(observed_y)
    search_lower = np.asarray(search_lower, dtype=np.float64)
    search_upper = np.asarray(search_upper, dtype=np.float64)
    if search_lower.shape != observed_y.shape or search_upper.shape != observed_y.shape:
        raise ValueError("search bounds must have one entry per observed y")
    if np.any(search_lower < 0) or np.any(search_upper < search_lower):
        raise ValueError("search bounds must satisfy 0 <= lower <= upper")
    if step <= 0:
        raise ValueError("grid step must be positive")
    if pair_chunk < 1:
        raise ValueError("pair_chunk must be positive")
    if not 0.0 < alpha < 1.0:
        raise ValueError("alpha must lie strictly between zero and one")

    if backend not in {"cpu", "torch"}:
        raise ValueError("backend must be 'cpu' or 'torch'")
    pair_y: list[np.ndarray] = []
    pair_lambda: list[np.ndarray] = []
    pair_ranges: list[tuple[int, int]] = []
    offset = 0
    for y, lower, upper in zip(observed_y, search_lower, search_upper):
        first = int(np.ceil(lower / step - 1e-9))
        last = int(np.floor(upper / step + 1e-9))
        if last < first:
            raise RuntimeError(f"search interval for y={int(y)} contains no grid point")
        candidates = np.round(np.arange(first, last + 1, dtype=np.float64) * step, 10)
        pair_lambda.append(candidates)
        pair_y.append(np.full(len(candidates), y, dtype=np.int64))
        pair_ranges.append((offset, offset + len(candidates)))
        offset += len(candidates)
    all_y = np.concatenate(pair_y)
    all_lambda = np.concatenate(pair_lambda)
    if backend == "cpu":
        h_pairs = h_pairs_cpu(
            all_y, all_lambda, table, background, quantile, grid_step=step
        )
    else:
        h_pairs = h_pairs_torch(
            all_y,
            all_lambda,
            table,
            background,
            quantile,
            device,
            pair_chunk=pair_chunk,
            grid_step=step,
        )

    lower = np.empty(len(observed_y), dtype=np.float64)
    upper = np.empty(len(observed_y), dtype=np.float64)
    for j, y in enumerate(observed_y):
        start, stop = pair_ranges[j]
        ix = np.flatnonzero(h_pairs[start:stop] > alpha)
        if len(ix) == 0:
            raise RuntimeError(f"no grid point satisfies h>alpha for y={int(y)}")
        lower[j] = all_lambda[start + ix[0]] + background
        upper[j] = all_lambda[start + ix[-1]] + background
    return IntervalTable(lower=lower, upper=upper)


def initial_search_bounds(
    method: str,
    initial_table: IntervalTable,
    observed_y: np.ndarray | list[int],
    background: float = 3.0,
) -> tuple[np.ndarray, np.ndarray]:
    """Use exact-start bounds and the documented RW/IM search-range rule.

    For y=0,1 the legacy script used U+5 as its upper search point; preserving
    that rule keeps the small-count outputs aligned with the existing tables.
    """

    ys = initial_table.require_y(observed_y)
    if method.upper() in EXACT_INITIAL_METHODS:
        return initial_table.lower[ys] - background, initial_table.upper[ys] - background
    if method.upper() in {"RW", "IM"}:
        # The text script used U+5 for y=0,1 before switching to y+3*sqrt(y).
        # Workbooks and the saved one-time intervals retain that small-y rule.
        upper = ys + 3.0 * np.sqrt(ys)
        small_y = ys <= 1
        upper[small_y] = initial_table.upper[ys[small_y]] - background + 5.0
        return np.zeros(len(ys), dtype=np.float64), upper
    raise ValueError(f"unknown method {method!r}; expected FC, RW, EB, or IM")


def _support_horizon(
    initial_table: IntervalTable,
    y_max: int,
    background: float,
    quantile: float,
) -> int:
    ys = np.arange(y_max + 1, dtype=np.int64)
    if y_max > initial_table.max_y:
        raise ValueError(f"initial table ends at y={initial_table.max_y}, requested y={y_max}")
    safe_upper = np.maximum(initial_table.upper[ys] - background, ys + 3.0 * np.sqrt(ys))
    small_y = ys <= 1
    safe_upper[small_y] = np.maximum(
        safe_upper[small_y], initial_table.upper[ys[small_y]] - background + 5.0
    )
    cutoff = poisson.ppf(quantile, float(np.max(safe_upper) + background))
    return max(y_max, int(cutoff))


def refinement_sequence(
    method: str,
    initial_table: IntervalTable,
    target_y_max: int,
    steps: int,
    alpha: float = 0.1,
    quantile: float = 0.999999,
    background: float = 3.0,
    grid_step: float = 0.01,
    backend: Backend = "torch",
    device: str | torch.device | None = None,
    pair_chunk: int = 16384,
) -> dict[int, IntervalTable]:
    """Compute iterations 1..steps for y<=target_y_max with support closure.

    Extra rows are calculated at early steps when later iterations need them as
    alternative outcomes. The required horizon is bounded using the initial
    upper endpoints and the manuscript's y+3*sqrt(y) search rule.
    """

    if target_y_max < 0 or steps < 1:
        raise ValueError("target_y_max must be nonnegative and steps must be positive")
    horizons = [int(target_y_max)]
    for _ in range(steps):
        horizons.append(_support_horizon(initial_table, horizons[-1], background, quantile))
    if horizons[-1] > initial_table.max_y:
        raise ValueError(
            f"initial interval data end at y={initial_table.max_y}, but {steps} refinements "
            f"need support through y={horizons[-1]}"
        )

    history: dict[int, IntervalTable] = {}
    current = initial_table
    for level in range(1, steps + 1):
        output_y_max = horizons[steps - level]
        ys = np.arange(output_y_max + 1, dtype=np.int64)
        if level == 1:
            search_lower, search_upper = initial_search_bounds(method, initial_table, ys, background)
        else:
            if output_y_max > current.max_y:
                raise RuntimeError("support closure did not provide enough previous-step intervals")
            search_lower = current.lower[ys] - background
            search_upper = current.upper[ys] - background
        current = refine_intervals(
            current,
            ys,
            search_lower,
            search_upper,
            alpha=alpha,
            quantile=quantile,
            background=background,
            step=grid_step,
            backend=backend,
            device=device,
            pair_chunk=pair_chunk,
        )
        history[level] = current
    return history


def available_refinement_sheets(path: str | Path) -> list[str]:
    """Return consecutive project iteration sheets present in a workbook."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        names = set(workbook.sheetnames)
    finally:
        workbook.close()
    sheets = ["Initial Interval"]
    level = 1
    while f"Refinement {level}" in names:
        sheets.append(f"Refinement {level}")
        level += 1
        if level > 8:
            break
    return sheets
